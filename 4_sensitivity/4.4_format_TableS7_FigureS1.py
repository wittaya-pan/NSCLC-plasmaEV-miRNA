# Builds Word versions of the Reviewer 1 sensitivity table and forest-plot figure
# from the CSVs written by 4.3_TableS7_FigureS1.R.
# Formatting only - no analysis. Run from project root: python 4_sensitivity/4.4_format_TableS7_FigureS1.py
import csv, os, re
import docx
from docx.shared import Pt, Mm
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

D = "results/R1_sensitivity_combined"
TNUM, FNUM = "S7", "S1"
EN, MINUS = "–", "−"

A = list(csv.DictReader(open(os.path.join(D, "R1_sensitivity_panelA_miRNAs.csv"), encoding="utf-8")))
B = list(csv.DictReader(open(os.path.join(D, "R1_sensitivity_panelB_riskscore.csv"), encoding="utf-8")))

def num(x, d=2):
    if x in ("", "NA", None): return "ND"
    v = float(x); s = ("%." + str(d) + "f") % abs(v)
    return (MINUS + s) if v < 0 else s
def pv(x):
    if x in ("", "NA", None): return "ND"
    v = float(x); return "<0.001" if v < 0.001 else "%.3f" % v
def hr(h, l, u):
    if h in ("", "NA", None): return "ND"
    return "%s (%s%s%s)" % (num(h), num(l), EN, num(u))

def runs(p, text, bold=False, size=10):
    for part in re.split(r"(\{[ab2]\})", text):
        if not part: continue
        sup = part in ("{a}", "{b}"); sub = part == "{2}"
        r = p.add_run(part[1] if (sup or sub) else part)
        r.bold = bold; r.font.size = Pt(size); r.font.name = "Times New Roman"
        if sup: r.font.superscript = True
        if sub: r.font.subscript = True

def base_doc(landscape):
    doc = docx.Document()
    s = doc.sections[0]
    if landscape:
        s.orientation = WD_ORIENT.LANDSCAPE; s.page_width, s.page_height = Mm(279.4), Mm(215.9)
    else:
        s.page_width, s.page_height = Mm(215.9), Mm(279.4)
    for m in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(s, m, Mm(20))
    st = doc.styles["Normal"]; st.font.name = "Times New Roman"; st.font.size = Pt(10)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    return doc

def para(doc, text, bold=False, size=10, align=WD_ALIGN_PARAGRAPH.LEFT, after=4):
    p = doc.add_paragraph(); p.alignment = align
    p.paragraph_format.space_after = Pt(after); p.paragraph_format.line_spacing = 1.0
    runs(p, text, bold, size)
    return p

def borders(table, top=True, bottom=True):
    tbl = table._tbl; pr = tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for edge, on in (("top", top), ("bottom", bottom), ("left", False), ("right", False),
                     ("insideH", False), ("insideV", False)):
        e = OxmlElement("w:" + edge)
        e.set(qn("w:val"), "single" if on else "nil"); e.set(qn("w:sz"), "8")
        e.set(qn("w:color"), "000000"); b.append(e)
    pr.append(b)

def underline_row(row):
    for c in row.cells:
        tcPr = c._tc.get_or_add_tcPr(); bd = OxmlElement("w:tcBorders")
        e = OxmlElement("w:bottom"); e.set(qn("w:val"), "single"); e.set(qn("w:sz"), "6")
        e.set(qn("w:color"), "000000"); bd.append(e); tcPr.append(bd)

def cell(c, text, bold=False, align=WD_ALIGN_PARAGRAPH.CENTER, size=9):
    c.text = ""; p = c.paragraphs[0]; p.alignment = align
    p.paragraph_format.space_after = Pt(1)
    runs(p, text, bold, size)

def make_table(doc, header, rows, widths, sections):
    t = doc.add_table(rows=1, cols=len(header)); borders(t)
    for i, h in enumerate(header):
        cell(t.rows[0].cells[i], h, bold=True,
             align=WD_ALIGN_PARAGRAPH.LEFT if i == 0 else WD_ALIGN_PARAGRAPH.CENTER)
    underline_row(t.rows[0])
    for kind, content in rows:
        r = t.add_row()
        if kind == "section":
            m = r.cells[0].merge(r.cells[-1]); cell(m, content, bold=True, align=WD_ALIGN_PARAGRAPH.LEFT)
        else:
            for i, v in enumerate(content):
                cell(r.cells[i], v, align=WD_ALIGN_PARAGRAPH.LEFT if i == 0 else WD_ALIGN_PARAGRAPH.CENTER)
    for row in t.rows:
        for i, w in enumerate(widths):
            row.cells[i].width = Mm(w)
    return t

# ------------------------------------------------------------------ table
doc = base_doc(landscape=True)
para(doc, "Table %s. Sensitivity of the three candidate miRNAs and the 3-miRNA risk score "
          "to outcome grouping and disease stage." % TNUM, bold=True, size=11, after=6)

para(doc, "A. Candidate miRNAs", bold=True, size=10, after=2)
grp = {"Primary: status at last follow-up": "34 (18 / 16)"}
order_mir = ["miR-377-3p", "miR-378h", "miR-509-3p"]
rowsA = []
for m in order_mir:
    rowsA.append(("section", m))
    for r in [x for x in A if x["miRNA"] == m]:
        g = r["n_groups"] if r["analysis"] != "Stage IV only" else "30"
        rowsA.append(("data", [r["analysis"], g, num(r["DE_log2FC"]), pv(r["DE_P"]),
                               hr(r["HR_perSD"], r["HR_lo"], r["HR_hi"]), pv(r["Cox_P"])]))
make_table(doc, ["Analysis", "Patients (deceased / surviving){a}", "DE log{2}FC", "DE P",
                 "HR per SD (95% CI)", "Cox P"], rowsA, [70, 50, 25, 22, 45, 22], None)

para(doc, "", after=4)
para(doc, "B. 3-miRNA risk score (fixed coefficients)", bold=True, size=10, after=2)
rowsB = []
for r in B:
    rowsB.append(("data", [r["analysis"], r["n"], r["events"], num(r["C_index"], 3),
                           hr(r["HR_perSD"], r["HR_lo"], r["HR_hi"]), pv(r["HR_P"]),
                           hr(r["HR_high_low"], r["HL_lo"], r["HL_hi"]), pv(r["logrank_P"])]))
make_table(doc, ["Analysis", "n", "Deaths", "C-index{b}", "HR per SD (95% CI)", "P",
                 "High vs low risk HR (95% CI)", "Log-rank P"], rowsB, [62, 12, 16, 20, 42, 18, 48, 20], None)

notes = [
 "{a}Deceased versus surviving at last follow-up (primary) or at 12 or 18 months (landmark; patients alive "
 "with shorter follow-up excluded). DE was not re-estimated in Stage IV only (Supplementary Methods).",
 "DE: limma after ComBat batch correction; log{2}FC, deceased versus surviving. Cox: univariate regression of "
 "overall survival; HR per 1 SD of z-scored expression.",
 "{b}Apparent C-index on batch-corrected expression, not optimism-corrected; the main-text value of 0.756 was "
 "computed without batch correction within the bootstrap. The risk score used the final model coefficients; "
 "high- and low-risk groups were defined by the full-cohort median (log-rank test).",
 "CI, confidence interval; C-index, concordance index; DE, differential expression; HR, hazard ratio; "
 "log{2}FC, log{2} fold change; ND, not done; SD, standard deviation.",
]
para(doc, "", after=2)
for n in notes: para(doc, n, size=9, after=3)
tpath = os.path.join(D, "Table %s.docx" % TNUM); doc.save(tpath)

# ------------------------------------------------------------------ figure
doc = base_doc(landscape=False)
para(doc, "Figure %s" % FNUM, bold=True, size=11, after=4)
doc.add_picture(os.path.join(D, "FigureS_R1_forest.png"), width=Mm(170))
legend = ("Figure %s. Hazard ratios of the three candidate miRNAs in sensitivity analyses. Univariate Cox "
  "regression of overall survival; squares, HR per 1 SD of z-scored expression; bars, 95%% CI; x-axis on log "
  "scale; dashed line, HR = 1. Primary: patients grouped for differential expression by survival status at last "
  "follow-up (n = 34, 18 deaths). 12- and 18-month landmark: grouped by survival status at 12 or 18 months, "
  "excluding patients alive with shorter follow-up (n = 28 and 25; 18 deaths each). Stage IV only: n = 30, 16 "
  "deaths. Values are given in Table S7. CI, confidence interval; HR, hazard ratio; SD, standard deviation."
  % FNUM)
para(doc, legend, size=10, after=4)
fpath = os.path.join(D, "Figure %s.docx" % FNUM); doc.save(fpath)
print("saved:", tpath, "|", fpath)
